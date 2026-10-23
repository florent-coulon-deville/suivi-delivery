"""Export de toute la configuration dans un fichier Excel (GET /api/configuration).
L'import du même fichier est dans test_import_configuration.py."""

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.configuration import FORMAT, SHEETS, TYPE_XLSX
from tests.conftest import create, installation_body


def export_file(client, headers: dict | None = None) -> dict[str, list[tuple]]:
    """GET /api/configuration qui doit réussir ; renvoie, par onglet, ses lignes (titres compris)."""
    response = client.get("/api/configuration", headers=headers or {})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}


def test_excel_response_to_download(client):
    response = client.get("/api/configuration")
    assert response.status_code == 200
    assert response.headers["content-type"] == TYPE_XLSX
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="suivi-delivery-configuration-')
    assert disposition.endswith('.xlsx"')


def test_empty_sheets_and_titles(client):
    sheets = export_file(client)
    assert list(sheets) == ["À propos", *SHEETS]
    for title, columns in SHEETS.items():
        assert sheets[title] == [columns]  # seulement la ligne des titres
    about = dict(sheets["À propos"])
    assert about["Format"] == FORMAT
    assert about["Version"] == client.app.version
    assert isinstance(about["Exporté le (UTC)"], datetime)


def test_objects_designated_by_their_codes(client, deployment):
    d = deployment
    other = create(client, "applications", identifier="socle", name="Socle", trigram="abc")
    auth = create(
        client, "modules", application_id=other["id"], code="auth", name="Authentification", type="CAPABILITY"
    )
    create(client, "module-versions", module_id=auth["id"], number="1.4", type="RELEASE")
    # Un module qui n'est installé nulle part : sa dépendance ne demande rien dans les environnements.
    ui = {"code": "ihm", "name": "Écrans", "type": "BUSINESS_PROJECT"}
    ui = create(client, "modules", application_id=d["application"]["id"], **ui)
    ui_1 = create(client, "module-versions", module_id=ui["id"], number="1.0")
    create(client, "dependencies", module_id=ui["id"], depends_on_id=auth["id"])
    create(
        client,
        "compatibilities",
        module_version_id=ui_1["id"],
        required_module_id=auth["id"],
        condition="AT_LEAST",
        version="1.4",
    )
    create(client, "installations", **installation_body(d, date="2026-10-01T08:30:00+02:00"))

    sheets = export_file(client)
    assert sheets["Environnements"][1:] == [("int", "Intégration", "non", None), ("prod", "Production", "oui", None)]
    assert sheets["Clusters"][1:] == [("cb01", "Cluster hors prod", "non"), ("cb02", "Cluster prod", "oui")]
    assert sheets["Zones réseau"][1:] == [("interne", "INTERNE"), ("partenaires", "PARTENAIRES")]
    assert sheets["Applications"][1:] == [("paiement", "Paiement", "abc"), ("socle", "Socle", "abc")]
    assert sheets["Namespaces"][1:] == [("paiement", "int", "cb01", "batch", "interne", "abc-paiement-int-batch")]
    assert sheets["Modules"][1:] == [
        ("paiement", "batch-paie", "Batch de paie", "BUSINESS_PROJECT"),
        ("paiement", "ihm", "Écrans", "BUSINESS_PROJECT"),
        ("socle", "auth", "Authentification", "CAPABILITY"),
    ]
    assert sheets["Dépendances"][1:] == [("paiement", "ihm", "socle", "auth")]
    assert sheets["Versions"][1:] == [
        ("paiement", "batch-paie", "2.3.0", "RELEASE"),
        ("paiement", "ihm", "1.0", "SNAPSHOT"),
        ("socle", "auth", "1.4", "RELEASE"),
    ]
    assert sheets["Compatibilités"][1:] == [("paiement", "ihm", "1.0", "socle", "auth", "AT_LEAST", "1.4", None)]
    assert sheets["Paramètres"][1:] == [
        ("paiement", "batch-paie", "APPLICATION", "appli"),
        ("paiement", "batch-paie", "TECHNICAL", "tech"),
    ]
    assert sheets["Affectations"][1:] == [
        ("paiement", "batch-paie", "APPLICATION", "appli", "int", "cb01", "batch"),
        ("paiement", "batch-paie", "TECHNICAL", "tech", "int", "cb01", "batch"),
    ]
    assert sheets["Emplacements"][1:] == [("paiement", "int", "cb01", "batch", "batch-paie", "batch-paie")]
    assert sheets["Installations"][1:] == [
        ("paiement", "int", "cb01", "batch", "batch-paie", "2.3.0", "tech", "appli", datetime(2026, 10, 1, 6, 30))
    ]


def test_environments_in_promotion_order(client):
    create(client, "environments", code="prod", label="Production", is_prod=True, promotion_order=3)
    create(client, "environments", code="int", label="Intégration", promotion_order=1)
    create(client, "environments", code="bac", label="Bac à sable")
    create(client, "environments", code="rec", label="Recette", promotion_order=2)
    codes = [row[0] for row in export_file(client)["Environnements"][1:]]
    assert codes == ["int", "rec", "prod", "bac"]  # sans ordre : à la fin


def test_methods(client):
    assert client.head("/api/configuration").status_code == 200
    response = client.options("/api/configuration")
    assert (response.status_code, response.headers["allow"]) == (204, "GET, HEAD, PATCH, OPTIONS")
    assert client.delete("/api/configuration").status_code == 405
