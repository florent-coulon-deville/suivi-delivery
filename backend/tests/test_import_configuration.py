"""Import de la configuration (PATCH /api/configuration) : le fichier Excel de l'export, en mode fusion."""

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.configuration import FORMAT, SHEETS, TYPE_XLSX
from tests.conftest import application, connected, create, installation_body, token

PASSWORD = "a-strong-password"


# --- Outils -------------------------------------------------------------------------


def export_file(client) -> bytes:
    response = client.get("/api/configuration")
    assert response.status_code == 200
    return response.content


def content(file: bytes) -> dict[str, list[tuple]]:
    """Les lignes de chaque onglet, sauf « À propos » (sa date change à chaque export)."""
    workbook = load_workbook(BytesIO(file))
    return {f.title: list(f.iter_rows(values_only=True)) for f in workbook.worksheets if f.title != "À propos"}


def file(format_: int | None = FORMAT, **sheets: list[tuple]) -> bytes:
    """Un fichier Excel fait à la main : les onglets donnés (titres des colonnes ajoutés), dans l'ordre
    de l'export. Onglet désigné par son nom anglais (network_zones=… pour « Zones réseau »), voir SHEET_NAMES."""
    workbook = Workbook()
    about = workbook.active
    assert about is not None
    about.title = "À propos"
    if format_ is not None:
        about.append(("Format", format_))
    names = dict(zip(SHEET_NAMES, SHEETS, strict=True))
    for key, rows in sheets.items():
        sheet = workbook.create_sheet(names.get(key, key))
        sheet.append(SHEETS.get(names.get(key, ""), ("Inconnu",)))
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


# Les onglets de SHEETS (titres en français), dans le même ordre, désignés en anglais dans les tests.
SHEET_NAMES = (
    "environments",
    "clusters",
    "network_zones",
    "applications",
    "namespaces",
    "modules",
    "dependencies",
    "versions",
    "compatibilities",
    "parameters",
    "assignments",
    "slots",
    "installations",
)


def import_workbook(client, content_: bytes, dry_run: bool = False, headers: dict | None = None):
    return client.patch(
        "/api/configuration",
        content=content_,
        params={"dry_run": "true"} if dry_run else {},
        headers={"Content-Type": TYPE_XLSX, **(headers or {})},
    )


def report(client, content_: bytes, dry_run: bool = False) -> dict:
    """Import (ou essai) qui doit réussir ; renvoie le rapport."""
    response = import_workbook(client, content_, dry_run)
    assert response.status_code == 200, response.json()
    return response.json()


def summary(report_: dict) -> dict[str, tuple[int, int, int]]:
    return {b["sheet"]: (b["additions"], b["updates"], b["unchanged"]) for b in report_["sheets"]}


def complete(client, d: dict) -> None:
    """Le jeu « deploiement », installé, avec une dépendance et une compatibilité vers un autre module."""
    create(client, "installations", **installation_body(d))  # date : maintenant, à la microseconde
    base = create(client, "applications", identifier="socle", name="Socle", trigram="soc")
    auth = create(client, "modules", application_id=base["id"], code="auth", name="Authentification", type="CAPABILITY")
    ui = {"code": "ihm", "name": "Écrans", "type": "BUSINESS_PROJECT"}
    ui = create(client, "modules", application_id=d["application"]["id"], **ui)
    ui_1 = create(client, "module-versions", module_id=ui["id"], number="1.0")
    create(client, "dependencies", module_id=ui["id"], depends_on_id=auth["id"])
    compatibility = {"condition": "BETWEEN", "version": "1.4", "version_max": "2.0"}
    create(client, "compatibilities", module_version_id=ui_1["id"], required_module_id=auth["id"], **compatibility)


# --- Aller-retour ---------------------------------------------------------------------


def test_reimporting_own_export_changes_nothing(client, deployment):
    complete(client, deployment)
    result = report(client, export_file(client), dry_run=True)
    assert result["problems"] == []
    assert result["changes"] == []
    assert all(additions == updates == 0 for additions, updates, _ in summary(result).values())
    assert summary(result)["Installations"] == (0, 0, 1)  # même date, à la seconde près


def test_copy_to_an_empty_instance(client, deployment):
    complete(client, deployment)
    export = export_file(client)
    with connected(application()) as other:  # une autre instance, vide
        result = report(other, export)
        assert result["applied"] is True
        assert summary(result)["Namespaces"] == (1, 0, 0)
        assert summary(result)["Affectations"] == (2, 0, 0)
        assert content(export_file(other)) == content(export)  # la même configuration


# --- Fusion -----------------------------------------------------------------------------


def test_dry_run_then_import(client, dataset):
    new = file(environments=[("int", "Intégration 2", "non", 1), ("rec", "Tests", "non", 2)])
    preview = report(client, new, dry_run=True)
    assert (preview["dry_run"], preview["applied"]) == (True, False)
    assert summary(preview) == {"Environnements": (1, 1, 0)}
    update = next(c for c in preview["changes"] if c["action"] == "UPDATE")
    assert (update["sheet"], update["row"], update["object"]) == ("Environnements", 2, "int")
    assert update["details"] == ["Libellé : Intégration → Intégration 2", "Ordre : (vide) → 1"]
    assert client.get("/api/environments").json()[0]["label"] == "Intégration"  # rien d'enregistré

    result = report(client, new)
    assert result["applied"] is True
    environments = {e["code"]: e for e in client.get("/api/environments").json()}
    assert environments["int"]["label"] == "Intégration 2"
    assert environments["rec"]["promotion_order"] == 2
    assert environments["prod"]["label"] == "Production"  # absent du fichier : pas touché


def test_import_logged_in_audit_log(client, dataset):
    report(client, file(clusters=[("cb03", "Cluster 3", "non")]))
    [row] = client.get("/api/audit-log", params={"object": "clusters", "action": "CREATE", "limit": 1}).json()
    assert (row["author"], row["after"]["code"]) == ("admin", "cb03")


def test_partial_file_referring_to_existing_objects(client, deployment):
    result = report(client, file(versions=[("PAIEMENT", "batch-paie", "2.4.0", "release")]))
    assert summary(result) == {"Versions": (1, 0, 0)}
    numbers = {v["number"]: v["type"] for v in client.get("/api/module-versions").json()}
    assert numbers == {"2.3.0": "RELEASE", "2.4.0": "RELEASE"}


def test_update_of_a_reference(client, deployment):
    row = ("paiement", "int", "cb01", "batch", "partenaires", "")
    result = report(client, file(namespaces=[row]))
    [change] = result["changes"]
    assert change["details"] == ["network_zone_id : interne → partenaires"]
    assert client.get("/api/namespaces").json()[0]["network_zone_id"] == deployment["zone_partenaires"]["id"]


def test_assignment_adds_a_namespace(client, deployment):
    d = deployment
    other = {"environment_id": d["int"]["id"], "cluster_id": d["cluster"]["id"], "network_zone_id": d["zone"]["id"]}
    create(client, "namespaces", application_id=d["application"]["id"], suffix="web", **other)
    assignments = [
        ("paiement", "batch-paie", "TECHNICAL", "tech", "int", "cb01", "batch"),  # déjà affecté
        ("paiement", "batch-paie", "TECHNICAL", "tech", "int", "cb01", "web"),
    ]
    assert summary(report(client, file(assignments=assignments))) == {"Affectations": (1, 0, 1)}
    technical_parameter = client.get(f"/api/parameters/{d['technical_parameter']['id']}").json()
    assert len(technical_parameter["namespace_ids"]) == 2


def test_dependency_and_installation_of_required_module_in_same_file(client, deployment):
    # batch-paie tourne en int. Déclarer qu'il a besoin de auth n'est accepté que si auth y tourne aussi :
    # écran par écran, il faudrait d'abord installer auth. Le fichier, lui, est jugé en entier.
    create(client, "installations", **installation_body(deployment))
    ns = ("socle", "int", "cb01", "app")
    result = report(
        client,
        file(
            applications=[("socle", "Socle", "soc")],
            namespaces=[(*ns, "interne", "")],
            modules=[("socle", "auth", "Authentification", "CAPABILITY")],
            dependencies=[("paiement", "batch-paie", "socle", "auth")],
            versions=[("socle", "auth", "1.4", "RELEASE")],
            parameters=[("socle", "auth", "TECHNICAL", "t1"), ("socle", "auth", "APPLICATION", "a1")],
            assignments=[
                ("socle", "auth", "TECHNICAL", "t1", "int", "cb01", "app"),
                ("socle", "auth", "APPLICATION", "a1", "int", "cb01", "app"),
            ],
            slots=[(*ns, "auth", "auth")],
            installations=[(*ns, "auth", "1.4", "t1", "a1", datetime(2026, 10, 1, 8, 0))],
        ),
    )
    assert result["problems"] == []
    assert len(client.get("/api/dependencies").json()) == 1


# --- Problèmes : rien n'est importé --------------------------------------------------------


def test_problems_with_sheet_and_row(client, dataset):
    new = file(
        clusters=[("cb03", "Cluster 3", "non")],  # valable, mais rien ne sera importé
        modules=[
            ("paiement", "batch", "Batch", "BUSINESS_PROJECT"),
            ("unknown", "web", "Web", "BUSINESS_PROJECT"),
            ("paiement", "batch", "Batch encore", "BUSINESS_PROJECT"),
            ("paiement", "api", "API", "WRONG_TYPE"),
        ],
    )
    preview = report(client, new, dry_run=True)
    assert [(p["sheet"], p["row"]) for p in preview["problems"]] == [
        ("Modules", 3),
        ("Modules", 4),
        ("Modules", 5),
    ]
    assert "application « unknown » introuvable" in preview["problems"][0]["message"]
    assert preview["problems"][1]["message"] == "même objet qu'à la ligne 2"
    assert preview["problems"][2]["message"].startswith("Type : ")

    response = import_workbook(client, new)
    assert response.status_code == 409
    assert response.json()["detail"] == "Import refusé : 3 problème(s), rien n'a été importé"
    assert response.json()["report"]["applied"] is False
    assert [c["code"] for c in client.get("/api/clusters").json()] == ["cb01", "cb02"]  # pas de cb03


def test_overall_rule_broken(client, deployment):
    ns = ("paiement", "prod", "cb02", "batch")
    new = file(
        namespaces=[(*ns, "interne", "")],
        versions=[("paiement", "batch-paie", "2.4.0-SNAPSHOT", "SNAPSHOT")],
        assignments=[
            ("paiement", "batch-paie", "TECHNICAL", "tech", "prod", "cb02", "batch"),
            ("paiement", "batch-paie", "APPLICATION", "appli", "prod", "cb02", "batch"),
        ],
        slots=[(*ns, "batch-paie", "batch-paie")],
        installations=[(*ns, "batch-paie", "2.4.0-SNAPSHOT", "tech", "appli", "2026-10-01T08:00:00")],
    )
    [problem] = report(client, new, dry_run=True)["problems"]
    assert (problem["sheet"], problem["row"]) == (None, None)
    assert "snapshot" in problem["message"]


def test_badly_filled_cells(client, dataset):
    new = file(
        environments=[("preprod", "Recette", "peut-être", 3), ("pprod", "", "non", "two")],
        installations=[("paiement", "int", "cb01", "batch", "x", "1.0", "t", "a", "yesterday")],
    )
    messages = [p["message"] for p in report(client, new, dry_run=True)["problems"]]
    assert messages == [
        "la colonne « Prod » doit valoir « oui » ou « non »",
        "la colonne « Libellé » est vide",
        "namespace « paiement / int / cb01 / batch » introuvable, ni dans l'outil ni dans le fichier",
    ]


# --- Fichier refusé en entier ------------------------------------------------------------------


def test_not_an_excel_file(client):
    response = import_workbook(client, b"pas un fichier excel")
    assert response.status_code == 400
    assert response.json()["detail"] == "Fichier illisible : ce n'est pas un fichier Excel (.xlsx)"


def test_content_type_required(client):
    response = client.patch("/api/configuration", json={"environments": []})
    assert response.status_code == 415
    assert "un fichier Excel" in response.json()["detail"]


def test_unknown_sheet(client):
    response = import_workbook(client, file(environmnts=[("int",)]))
    assert response.status_code == 422
    assert response.json()["detail"].startswith("Onglet inconnu : « environmnts »")


def test_unexpected_columns(client):
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Clusters"
    sheet.append(("Code", "Prod", "Libellé"))
    output = BytesIO()
    workbook.save(output)
    response = import_workbook(client, output.getvalue())
    assert response.status_code == 422
    assert response.json()["detail"] == "Onglet Clusters : colonnes attendues, dans l'ordre : Code, Libellé, Prod"


def test_other_format(client):
    response = import_workbook(client, file(format_=FORMAT + 1))
    assert response.status_code == 422
    assert response.json()["detail"] == f"Fichier au format {FORMAT + 1} : cet outil lit le format {FORMAT}"


def test_reserved_to_admin(client):
    body = {"identifier": "paul", "name": "Paul", "kind": "PERSON", "role": "EDITOR"}
    assert client.post("/api/accounts", json={**body, "password": PASSWORD}).status_code == 201
    headers = {"Authorization": f"Bearer {token(client.app, 'paul')}"}
    assert import_workbook(client, file(), headers=headers).status_code == 403


def test_date_without_timezone_read_as_utc(client, deployment):
    row = ("paiement", "int", "cb01", "batch", "batch-paie", "2.3.0", "tech", "appli", "2026-10-01T08:30:00")
    report(client, file(installations=[row]))
    [installation] = client.get("/api/installations").json()
    assert datetime.fromisoformat(installation["date"]) == datetime(2026, 10, 1, 8, 30, tzinfo=UTC)
