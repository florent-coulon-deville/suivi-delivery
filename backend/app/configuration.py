"""Toute la configuration dans un fichier Excel : l'export (l'import est dans app/import_configuration.py).

Le fichier sert à trois choses : recopier la configuration d'une instance à l'autre, la
sauvegarder, et la saisir en masse dans un tableur.

Forme du fichier :
  - un onglet « À propos » (outil, version, format, date de l'export) ;
  - puis un onglet par objet, dans l'ordre de saisie (SHEETS) : un objet ne désigne que des
    objets des onglets précédents ;
  - première ligne : les titres des colonnes ; ensuite, une ligne par objet.

Les objets se désignent par leurs **noms métier**, jamais par les identifiants de la base (1, 2,
3…), qui changent d'une instance à l'autre. Exemples :
  - un module           : application + code du module ;
  - un namespace        : application + environnement + cluster + suffixe ;
  - un emplacement      : son namespace + nom de release ;
  - un paramètre        : module + type + nom.
Quand le lien est forcément dans la même application (le module d'un emplacement, les paramètres
d'une installation…), on n'écrit que le code ou le nom.

Valeurs : oui / non pour les cases à cocher (prod), dates en UTC, le reste tel que dans l'outil
(en minuscules). Le nom calculé des namespaces est donné pour information.
"""

from collections.abc import Callable
from datetime import UTC, datetime
from enum import StrEnum
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.routes.namespaces import with_name
from app.storage import Storage, Table

# Numéro du format du fichier : à augmenter si les onglets ou les colonnes changent, pour que
# l'import reconnaisse un fichier d'un ancien format.
FORMAT = 1

TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Titre de la colonne du libellé (environnements, clusters, zones réseau), relu par l'import.
LABEL = "Libellé"

# Les onglets, dans l'ordre de saisie, et leurs colonnes.
SHEETS: dict[str, tuple[str, ...]] = {
    "Environnements": ("Code", LABEL, "Prod", "Ordre"),
    "Clusters": ("Code", LABEL, "Prod"),
    "Zones réseau": ("Code", LABEL),
    "Applications": ("Identifiant", "Nom", "Trigramme"),
    "Namespaces": ("Application", "Environnement", "Cluster", "Suffixe", "Zone réseau", "Nom (calculé)"),
    "Modules": ("Application", "Code", "Nom", "Type"),
    "Dépendances": ("Application", "Module", "Application requise", "Module requis"),
    "Versions": ("Application", "Module", "Numéro", "Type"),
    "Compatibilités": (
        "Application",
        "Module",
        "Version",
        "Application requise",
        "Module requis",
        "Condition",
        "Version requise",
        "Version max",
    ),
    "Paramètres": ("Application", "Module", "Type", "Nom"),
    # Les namespaces auxquels chaque paramètre est affecté : une ligne par paramètre et par namespace.
    "Affectations": ("Application", "Module", "Type", "Paramètre", "Environnement", "Cluster", "Suffixe"),
    "Emplacements": ("Application", "Environnement", "Cluster", "Suffixe", "Release", "Module"),
    "Installations": (
        "Application",
        "Environnement",
        "Cluster",
        "Suffixe",
        "Release",
        "Version",
        "Paramètre technique",
        "Paramètre applicatif",
        "Date (UTC)",
    ),
}


def file_name(now: datetime) -> str:
    """Nom proposé au téléchargement, ex. suivi-delivery-configuration-2026-10-14.xlsx."""
    return f"suivi-delivery-configuration-{now:%Y-%m-%d}.xlsx"


def excel_file(storage: Storage, version: str, now: datetime) -> bytes:
    """Le fichier Excel de toute la configuration (contenu du fichier .xlsx)."""
    workbook = Workbook()
    about = workbook.active
    assert about is not None  # un classeur neuf a toujours une feuille
    about.title = "À propos"
    for row in (
        ("Outil", "Suivi de la delivery"),
        ("Version", version),
        ("Format", FORMAT),
        ("Exporté le (UTC)", _to_utc(now)),
        ("Contenu", "Un onglet par objet, dans l'ordre de saisie ; les objets sont désignés par leurs codes."),
    ):
        about.append(row)
    _format(about, titles=False)

    for title, rows in rows_by_sheet(storage).items():
        sheet = workbook.create_sheet(title)
        sheet.append(SHEETS[title])
        for row in rows:
            sheet.append([_cell(value) for value in row])
        _format(sheet, titles=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def rows_by_sheet(storage: Storage) -> dict[str, list[list]]:
    """Le contenu de chaque onglet (sans la ligne des titres), objets désignés par leurs noms métier."""
    applications = _by_id(storage.applications)
    environments = _by_id(storage.environments)
    clusters = _by_id(storage.clusters)
    zones = _by_id(storage.network_zones)
    namespaces = _by_id(storage.namespaces)
    modules = _by_id(storage.modules)
    versions = _by_id(storage.module_versions)
    parameters = _by_id(storage.parameters)
    slots = _by_id(storage.slots)

    def application(id_: int) -> str:
        return applications[id_]["identifier"]

    def module(id_: int) -> list:
        """Application + code du module."""
        return [application(modules[id_]["application_id"]), modules[id_]["code"]]

    def namespace(id_: int) -> list:
        """Application + environnement + cluster + suffixe."""
        n = namespaces[id_]
        return [
            application(n["application_id"]),
            environments[n["environment_id"]]["code"],
            clusters[n["cluster_id"]]["code"],
            n["suffix"],
        ]

    def slot(id_: int) -> list:
        """Son namespace + nom de release."""
        return [*namespace(slots[id_]["namespace_id"]), slots[id_]["release_name"]]

    def parameter(id_: int) -> list:
        """Module + type + nom."""
        p = parameters[id_]
        return [*module(p["module_id"]), p["type"], p["name"]]

    content = {
        "Environnements": [
            [e["code"], e["label"], e["is_prod"], e["promotion_order"]]
            for e in sorted(environments.values(), key=_promotion_order)
        ],
        "Clusters": [[c["code"], c["label"], c["is_prod"]] for c in clusters.values()],
        "Zones réseau": [[z["code"], z["label"]] for z in zones.values()],
        "Applications": [[a["identifier"], a["name"], a["trigram"]] for a in applications.values()],
        "Namespaces": [
            [
                *namespace(n["id"]),
                zones[n["network_zone_id"]]["code"],
                with_name(storage, n)["name"],
            ]
            for n in namespaces.values()
        ],
        "Modules": [[*module(m["id"]), m["name"], m["type"]] for m in modules.values()],
        "Dépendances": [
            [*module(d["module_id"]), *module(d["depends_on_id"])] for d in storage.dependencies.list_all()
        ],
        "Versions": [[*module(v["module_id"]), v["number"], v["type"]] for v in versions.values()],
        "Compatibilités": [
            [
                *module(versions[c["module_version_id"]]["module_id"]),
                versions[c["module_version_id"]]["number"],
                *module(c["required_module_id"]),
                c["condition"],
                c["version"],
                c["version_max"],
            ]
            for c in storage.compatibilities.list_all()
        ],
        "Paramètres": [parameter(p["id"]) for p in parameters.values()],
        # Le namespace est dans l'application du paramètre : environnement, cluster, suffixe suffisent.
        "Affectations": [
            [*parameter(p["id"]), *namespace(namespace_id)[1:]]
            for p in parameters.values()
            for namespace_id in p["namespace_ids"]
        ],
        # Le module est dans l'application du namespace : son code suffit.
        "Emplacements": [[*slot(e["id"]), modules[e["module_id"]]["code"]] for e in slots.values()],
        # La version et les paramètres sont ceux du module de l'emplacement : numéro et noms suffisent.
        "Installations": [
            [
                *slot(i["slot_id"]),
                versions[i["module_version_id"]]["number"],
                parameters[i["technical_parameter_id"]]["name"],
                parameters[i["application_parameter_id"]]["name"],
                i["date"],
            ]
            for i in sorted(storage.installations.list_all(), key=lambda i: (i["date"], i["id"]))
        ],
    }
    # Toujours le même ordre, d'un export à l'autre (sauf environnements et installations, déjà triés).
    already_sorted = {"Environnements", "Installations"}
    return {title: rows if title in already_sorted else sorted(rows, key=_sort_key) for title, rows in content.items()}


def _by_id(table: Table) -> dict[int, dict]:
    return {row["id"]: row for row in table.list_all()}


def _promotion_order(environment: dict) -> tuple:
    """Les environnements dans l'ordre de promotion, ceux sans ordre à la fin, puis par code."""
    return (environment["promotion_order"] is None, environment["promotion_order"] or 0, environment["code"])


def _sort_key(row: list) -> list[str]:
    return ["" if value is None else str(value) for value in row]


def _to_utc(instant: datetime) -> datetime:
    """Excel ne connaît pas les fuseaux horaires : la date est écrite en UTC, sans fuseau, à la seconde
    près (l'import compare les installations à la seconde près). Sans fuseau dans l'outil : déjà en UTC."""
    utc = instant.astimezone(UTC) if instant.tzinfo is not None else instant
    return utc.replace(tzinfo=None, microsecond=0)


def _cell(value):
    """La valeur telle qu'écrite dans une case : oui / non, dates en UTC, énumérations en texte."""
    conversions: list[tuple[type, Callable]] = [
        (bool, lambda v: "oui" if v else "non"),
        (datetime, _to_utc),
        (StrEnum, lambda v: v.value),
    ]
    for type_, conversion in conversions:
        if isinstance(value, type_):
            return conversion(value)
    return value


def _format(sheet, titles: bool) -> None:
    """Titres en gras et figés, filtres, largeur des colonnes, format des dates."""
    if titles:
        for case in sheet[1]:
            case.font = Font(bold=True)
        sheet.freeze_panes = "A2"  # la ligne des titres reste visible
        sheet.auto_filter.ref = sheet.dimensions  # un filtre sur chaque colonne
    else:
        for (case,) in sheet.iter_rows(max_col=1):
            case.font = Font(bold=True)
    for column in sheet.iter_cols():
        for case in column:
            if isinstance(case.value, datetime):
                case.number_format = "yyyy-mm-dd hh:mm:ss"
        width = max(len(str(case.value)) if case.value is not None else 0 for case in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width, 8) + 2, 60)
