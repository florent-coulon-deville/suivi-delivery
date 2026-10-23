"""Import de la configuration : relit le fichier Excel de l'export (app/configuration.py).

Mode **fusion** (PATCH /api/configuration) : chaque ligne du fichier désigne un objet par ses noms
métier (codes), comme à l'export.
  - objet inconnu de l'outil   → ajouté ;
  - objet connu, autre contenu → mis à jour (libellé, nom, type, zone réseau…) ;
  - objet connu, même contenu  → inchangé ;
  - objet absent du fichier    → pas touché : un fichier peut ne contenir qu'une partie
                                 (un onglet seul, quelques lignes), qui désigne des objets
                                 déjà dans l'outil.
Les installations sont un historique : une installation déjà présente (même emplacement, version,
paramètres et date, à la seconde près) est inchangée, sinon ajoutée ; jamais modifiée.
Les affectations ajoutent un namespace à un paramètre ; elles n'en retirent jamais.

Règle d'or : **tout ou rien**. Les onglets sont lus dans l'ordre de saisie et chaque ligne passe
par les mêmes vérifications que la saisie écran par écran (les fonctions ajouter / enregistrer
des routes). Une ligne refusée est notée comme problème, avec son onglet et son numéro, et la
lecture continue pour tout signaler d'un coup. Les règles d'ensemble (compatibilités, pas de
snapshot en prod…) sont vérifiées une fois, à la fin, sur le résultat (deferred_consistency).
S'il reste un problème, ou pour un simple essai (aperçu), tout est annulé : l'import se fait dans
un essai du stockage (Storage.attempt), comme les saisies de app/consistency.py.

Chaque objet ajouté ou modifié est noté au journal, au nom de la personne qui importe.
"""

from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import ModuleType
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from pydantic_core import ErrorDetails

from app.configuration import FORMAT, LABEL, SHEETS
from app.consistency import deferred_consistency, inconsistencies
from app.routes import (
    applications,
    clusters,
    compatibilities,
    dependencies,
    environments,
    installations,
    module_versions,
    modules,
    namespaces,
    network_zones,
    parameters,
    slots,
)
from app.schemas import (
    ApplicationCreation,
    ClusterCreation,
    CompatibilityCreation,
    DependencyCreation,
    EnvironmentCreation,
    ImportAction,
    ImportChange,
    ImportProblem,
    ImportReport,
    InstallationCreation,
    ModuleCreation,
    ModuleVersionCreation,
    NamespaceCreation,
    NetworkZoneCreation,
    ParameterCreation,
    SheetSummary,
    SlotCreation,
)
from app.storage import Storage, Table

ABOUT = "À propos"

# Noms des champs dans les messages de validation (« libelle » → « Libellé »).
FIELDS = {
    "code": "Code",
    "label": LABEL,
    "promotion_order": "Ordre",
    "identifier": "Identifiant",
    "name": "Nom",
    "trigram": "Trigramme",
    "suffix": "Suffixe",
    "type": "Type",
    "number": "Numéro",
    "condition": "Condition",
    "version": "Version requise",
    "version_max": "Version max",
    "release_name": "Release",
    "date": "Date (UTC)",
}


class UnreadableFile(Exception):
    """Le contenu reçu n'est pas un fichier Excel (.xlsx) : 400."""


class InvalidFormat(Exception):
    """C'est un fichier Excel, mais pas celui de l'export (onglet ou colonnes inconnus, autre format) : 422."""


class RowProblem(Exception):
    """Une ligne du fichier est refusée : son message va au rapport, la lecture continue."""


class _Cancel(Exception):
    """Sort de l'essai du stockage pour tout annuler (aperçu, ou problèmes)."""


# --- Lecture du fichier ---------------------------------------------------------


type Sheets = dict[str, list[tuple[int, dict[str, Any]]]]


def read_file(content: bytes) -> Sheets:
    """Les lignes de chaque onglet du fichier : [(numéro de ligne, {colonne: valeur})].
    Lève UnreadableFile ou InvalidFormat."""
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception as error:  # zip abîmé, autre type de fichier… : openpyxl lève des erreurs variées
        raise UnreadableFile("Fichier illisible : ce n'est pas un fichier Excel (.xlsx)") from error
    sheets: Sheets = {}
    for sheet in workbook.worksheets:
        title = sheet.title.strip()
        rows = [list(values) for values in sheet.iter_rows(values_only=True)]
        if title == ABOUT:
            _check_format(rows)
            continue
        if title not in SHEETS:
            expected = ", ".join(SHEETS)
            raise InvalidFormat(f"Onglet inconnu : « {title} ». Onglets possibles : {ABOUT}, {expected}")
        if not rows:
            continue
        columns = SHEETS[title]
        titles = [_raw_text(value) for value in rows[0]]
        while titles and not titles[-1]:
            titles.pop()  # cases vides à droite des titres
        if titles != list(columns):
            raise InvalidFormat(f"Onglet {title} : colonnes attendues, dans l'ordre : {', '.join(columns)}")
        sheets[title] = [
            (number, dict(zip(columns, values + [None] * len(columns), strict=False)))
            for number, values in enumerate(rows[1:], start=2)
            if any(_raw_text(value) for value in values)  # lignes vides ignorées
        ]
    return sheets


def _check_format(rows: list[list]) -> None:
    """Onglet « À propos » : le numéro de format doit être celui de cet outil."""
    for row in rows:
        if row and _raw_text(row[0]) == "Format":
            format_ = row[1] if len(row) > 1 else None
            if format_ != FORMAT:
                raise InvalidFormat(f"Fichier au format {format_} : cet outil lit le format {FORMAT}")


def _raw_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


class Row:
    """Les cases d'une ligne du fichier, lues et contrôlées (RowProblem si une case est fausse)."""

    def __init__(self, values: dict[str, Any]) -> None:
        self._values = values

    def text(self, column: str) -> str:
        """Une case obligatoire, en texte (1.0 saisi comme nombre → « 1 »)."""
        text = self.optional_text(column)
        if text is None:
            raise RowProblem(f"la colonne « {column} » est vide")
        return text

    def optional_text(self, column: str) -> str | None:
        value = self._values.get(column)
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        text = _raw_text(value)
        return text or None

    def code(self, column: str) -> str:
        """Un code, un identifiant, un nom : toujours en minuscules dans l'outil."""
        return self.text(column).lower()

    def list(self, column: str) -> str:
        """Une valeur de liste fermée (type, condition) : en majuscules dans l'outil."""
        return self.text(column).upper()

    def yes_no(self, column: str) -> bool:
        value = self._values.get(column)
        if isinstance(value, bool):
            return value
        text = (self.optional_text(column) or "non").lower()
        if text in ("oui", "true", "true", "1"):
            return True
        if text in ("non", "false", "false", "0"):
            return False
        raise RowProblem(f"la colonne « {column} » doit valoir « oui » ou « non »")

    def integer(self, column: str) -> int | None:
        text = self.optional_text(column)
        if text is None:
            return None
        if not text.lstrip("-").isdigit():
            raise RowProblem(f"la colonne « {column} » doit être un nombre entier")
        return int(text)

    def date(self, column: str) -> datetime:
        """Une date en UTC, à la seconde près (une date Excel peut revenir à la milliseconde près)."""
        value = self._values.get(column)
        if isinstance(value, str):
            try:
                value = datetime.fromisoformat(value.strip())
            except ValueError:
                value = None
        if not isinstance(value, datetime):
            raise RowProblem(f"la colonne « {column} » doit être une date (ex. 2026-10-14 08:30:00)")
        value = value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
        return to_the_second(value + timedelta(microseconds=500_000))  # arrondie


def to_the_second(instant: datetime) -> datetime:
    return instant.replace(microsecond=0)


def to_utc(instant: datetime) -> datetime:
    """Une date de l'outil en UTC, à la seconde près (sans fuseau : déjà en UTC)."""
    return to_the_second(instant.replace(tzinfo=UTC) if instant.tzinfo is None else instant.astimezone(UTC))


# --- Import -----------------------------------------------------------------------


def import_workbook(storage: Storage, content: bytes, dry_run: bool) -> ImportReport:
    """Importe le fichier (mode fusion) ; dry_run=True : aperçu, rien n'est enregistré.
    S'il y a des problèmes, rien n'est enregistré non plus. Lève UnreadableFile ou InvalidFormat."""
    sheets = read_file(content)
    import_run = _ImportRun()
    try:
        with storage.attempt() as clone:
            import_run.storage = clone
            with deferred_consistency():
                for title, process in _PROCESSING.items():
                    for number, values in sheets.get(title, []):
                        import_run.row(title, number, values, process)
            # Les règles d'ensemble, une fois, sur le résultat.
            import_run.problems += [ImportProblem(sheet=None, row=None, message=m) for m in inconsistencies(clone)]
            if dry_run or import_run.problems:
                raise _Cancel
    except _Cancel:
        applied = False
    else:
        applied = True
    return ImportReport(
        dry_run=dry_run,
        applied=applied,
        sheets=[summary for title, summary in import_run.summaries.items() if title in sheets],
        changes=import_run.changes,
        problems=import_run.problems,
    )


class _ImportRun:
    """L'import en cours : le stockage (l'essai), et ce qu'on a fait de chaque ligne."""

    def __init__(self) -> None:
        self.storage: Storage
        self.summaries = {title: SheetSummary(sheet=title) for title in SHEETS}
        self.changes: list[ImportChange] = []
        self.problems: list[ImportProblem] = []
        self._views: dict[str, dict[tuple, int]] = {title: {} for title in SHEETS}  # objets déjà lus → ligne

    def row(self, sheet: str, number: int, values: dict, process: Callable) -> None:
        """Traite une ligne ; si elle est refusée, note le problème et continue."""
        self._sheet, self._number = sheet, number
        try:
            process(self, Row(values))
        except RowProblem as problem:
            self._problem(str(problem))
        except HTTPException as rejection:  # refus des vérifications des routes (404, 409, 422)
            self._problem(str(rejection.detail))
        except ValidationError as error:  # valeur au mauvais format (schémas de app/schemas.py)
            self._problem(" ; ".join(_validation_message(e) for e in error.errors()))

    def _problem(self, message: str) -> None:
        self.problems.append(ImportProblem(sheet=self._sheet, row=self._number, message=message))

    def once(self, key: tuple) -> None:
        """Un même objet ne doit apparaître qu'une fois dans l'onglet."""
        already = self._views[self._sheet].setdefault(key, self._number)
        if already != self._number:
            raise RowProblem(f"même objet qu'à la ligne {already}")

    def log(self, action: ImportAction | None, described: str, details: list[str] | None = None) -> None:
        summary = self.summaries[self._sheet]
        if action is None:
            summary.unchanged += 1
            return
        if action == ImportAction.ADD:
            summary.additions += 1
        else:
            summary.updates += 1
        change = ImportChange(
            sheet=self._sheet, row=self._number, action=action, object=described, details=details or []
        )
        self.changes.append(change)

    def merge(
        self,
        table: Table,
        route: ModuleType,
        key: dict,
        data: BaseModel,
        described: str,
        compare: set[str] | None = None,
    ) -> None:
        """Ajoute l'objet s'il n'existe pas (cle), le met à jour s'il a changé, sinon ne fait rien.
        route : le module de la route, dont on appelle ajouter / enregistrer (mêmes vérifications que l'écran).
        comparer : les champs à comparer (par défaut : tous ceux de donnees)."""
        self.once(tuple(key.values()))
        new = data.model_dump()
        existing = table.find(**key)
        if existing is None:
            route.add(self.storage, new)
            self.log(ImportAction.ADD, described)
            return
        fields = compare if compare is not None else set(new)
        changes = {field: new[field] for field in fields if existing[field] != new[field]}
        if not changes:
            self.log(None, described)
            return
        route.save(self.storage, existing["id"], changes)
        details = [
            f"{FIELDS.get(field, field)} : {self._readable(field, existing[field])} → {self._readable(field, value)}"
            for field, value in sorted(changes.items())
        ]
        self.log(ImportAction.UPDATE, described, details)

    def _readable(self, field: str, value: Any) -> str:
        """Une valeur telle qu'on la lit dans le fichier : codes plutôt qu'identifiants, oui / non…"""
        if field == "network_zone_id":
            return self.storage.network_zones.read_existing(value)["code"]
        if field == "module_id":
            return self.storage.modules.read_existing(value)["code"]
        if isinstance(value, bool):
            return "oui" if value else "non"
        return "(vide)" if value is None else str(value)

    # --- Retrouver les objets par leurs codes (dans l'outil, ou ajoutés par les lignes précédentes) ---

    def find(self, table: Table, description: str, **criteria) -> dict:
        row = table.find(**criteria)
        if row is None:
            raise RowProblem(f"{description} introuvable, ni dans l'outil ni dans le fichier")
        return row

    def application(self, identifier: str) -> dict:
        return self.find(self.storage.applications, f"application « {identifier} »", identifier=identifier)

    def module(self, application: str, code: str) -> dict:
        application_id = self.application(application)["id"]
        return self.find(
            self.storage.modules, f"module « {application} / {code} »", application_id=application_id, code=code
        )

    def namespace(self, cases: Row) -> dict:
        """Le namespace des colonnes Application, Environnement, Cluster, Suffixe."""
        application, environment, cluster, suffix = (
            cases.code(c) for c in ("Application", "Environnement", "Cluster", "Suffixe")
        )
        return self.find(
            self.storage.namespaces,
            f"namespace « {application} / {environment} / {cluster} / {suffix} »",
            application_id=self.application(application)["id"],
            environment_id=self.find(self.storage.environments, f"environnement « {environment} »", code=environment)[
                "id"
            ],
            cluster_id=self.find(self.storage.clusters, f"cluster « {cluster} »", code=cluster)["id"],
            suffix=suffix,
        )

    def version(self, module: dict, number: str) -> dict:
        return self.find(
            self.storage.module_versions,
            f"version « {module['code']} {number} »",
            module_id=module["id"],
            number=number,
        )

    def parameter(self, module: dict, type_: str, name: str) -> dict:
        return self.find(
            self.storage.parameters,
            f"paramètre {type_} « {module['code']} / {name} »",
            module_id=module["id"],
            type=type_,
            name=name,
        )


def _validation_message(error: ErrorDetails) -> str:
    field = ".".join(str(part) for part in error["loc"])
    return f"{FIELDS.get(field, field)} : {error['msg']}"


def _namespace_name(cases: Row) -> str:
    return " / ".join(cases.code(c) for c in ("Application", "Environnement", "Cluster", "Suffixe"))


# --- Un traitement par onglet, dans l'ordre de saisie -------------------------------


def _environment(i: _ImportRun, cases: Row) -> None:
    code = cases.code("Code")
    data = EnvironmentCreation(
        code=code, label=cases.text(LABEL), is_prod=cases.yes_no("Prod"), promotion_order=cases.integer("Ordre")
    )
    i.merge(i.storage.environments, environments, {"code": code}, data, code)


def _cluster(i: _ImportRun, cases: Row) -> None:
    code = cases.code("Code")
    data = ClusterCreation(code=code, label=cases.text(LABEL), is_prod=cases.yes_no("Prod"))
    i.merge(i.storage.clusters, clusters, {"code": code}, data, code)


def _network_zone(i: _ImportRun, cases: Row) -> None:
    code = cases.code("Code")
    data = NetworkZoneCreation(code=code, label=cases.text(LABEL))
    i.merge(i.storage.network_zones, network_zones, {"code": code}, data, code)


def _application(i: _ImportRun, cases: Row) -> None:
    identifier = cases.code("Identifiant")
    data = ApplicationCreation(identifier=identifier, name=cases.text("Nom"), trigram=cases.code("Trigramme"))
    i.merge(i.storage.applications, applications, {"identifier": identifier}, data, identifier)


def _namespace(i: _ImportRun, cases: Row) -> None:
    # « Nom (calculé) » n'est pas lu : l'outil le calcule.
    application, environment, cluster, suffix = (
        cases.code(c) for c in ("Application", "Environnement", "Cluster", "Suffixe")
    )
    zone = cases.code("Zone réseau")
    data = NamespaceCreation(
        application_id=i.application(application)["id"],
        environment_id=i.find(i.storage.environments, f"environnement « {environment} »", code=environment)["id"],
        cluster_id=i.find(i.storage.clusters, f"cluster « {cluster} »", code=cluster)["id"],
        suffix=suffix,
        network_zone_id=i.find(i.storage.network_zones, f"zone réseau « {zone} »", code=zone)["id"],
    )
    key = {
        "application_id": data.application_id,
        "environment_id": data.environment_id,
        "cluster_id": data.cluster_id,
        "suffix": data.suffix,
    }
    i.merge(i.storage.namespaces, namespaces, key, data, _namespace_name(cases))


def _module(i: _ImportRun, cases: Row) -> None:
    application, code = cases.code("Application"), cases.code("Code")
    data = ModuleCreation(
        application_id=i.application(application)["id"], code=code, name=cases.text("Nom"), type=cases.list("Type")
    )
    key = {"application_id": data.application_id, "code": code}
    i.merge(i.storage.modules, modules, key, data, f"{application} / {code}")


def _dependency(i: _ImportRun, cases: Row) -> None:
    module = i.module(cases.code("Application"), cases.code("Module"))
    required = i.module(cases.code("Application requise"), cases.code("Module requis"))
    data = DependencyCreation(module_id=module["id"], depends_on_id=required["id"])
    described = (
        f"{cases.code('Application')} / {module['code']} → {cases.code('Application requise')} / {required['code']}"
    )
    i.merge(i.storage.dependencies, dependencies, data.model_dump(), data, described)


def _version(i: _ImportRun, cases: Row) -> None:
    application = cases.code("Application")
    module = i.module(application, cases.code("Module"))
    number = cases.text("Numéro")
    data = ModuleVersionCreation(module_id=module["id"], number=number, type=cases.list("Type"))
    key = {"module_id": module["id"], "number": number}
    described = f"{application} / {module['code']} {number}"
    i.merge(i.storage.module_versions, module_versions, key, data, described)


def _compatibility(i: _ImportRun, cases: Row) -> None:
    application = cases.code("Application")
    module = i.module(application, cases.code("Module"))
    version = i.version(module, cases.text("Version"))
    required = i.module(cases.code("Application requise"), cases.code("Module requis"))
    data = CompatibilityCreation(
        module_version_id=version["id"],
        required_module_id=required["id"],
        condition=cases.list("Condition"),
        version=cases.text("Version requise"),
        version_max=cases.optional_text("Version max"),
    )
    key = {"module_version_id": version["id"], "required_module_id": required["id"]}
    described = f"{application} / {module['code']} {version['number']} → {required['code']}"
    i.merge(i.storage.compatibilities, compatibilities, key, data, described)


def _parameter(i: _ImportRun, cases: Row) -> None:
    application = cases.code("Application")
    module = i.module(application, cases.code("Module"))
    type_, name = cases.list("Type"), cases.code("Nom")
    # Créé sans namespace : ses namespaces viennent de l'onglet Affectations.
    data = ParameterCreation(module_id=module["id"], name=name, type=type_, namespace_ids=[])
    key = {"module_id": module["id"], "type": data.type, "name": data.name}
    described = f"{application} / {module['code']} {type_} {name}"
    # Rien d'autre à comparer que la clé : un paramètre existant reste tel quel (namespaces compris).
    i.merge(i.storage.parameters, parameters, key, data, described, compare=set())


def _assignment(i: _ImportRun, cases: Row) -> None:
    application = cases.code("Application")
    module = i.module(application, cases.code("Module"))
    type_, name = cases.list("Type"), cases.code("Paramètre")
    parameter = i.parameter(module, type_, name)
    namespace = i.namespace(cases)
    i.once((parameter["id"], namespace["id"]))
    described = f"{application} / {module['code']} {type_} {name} → {_namespace_name(cases)}"
    if namespace["id"] in parameter["namespace_ids"]:
        i.log(None, described)
        return
    assigned = [*parameter["namespace_ids"], namespace["id"]]
    parameters.save(i.storage, parameter["id"], {"namespace_ids": assigned})
    i.log(ImportAction.ADD, described)


def _slot(i: _ImportRun, cases: Row) -> None:
    namespace = i.namespace(cases)
    module = i.module(cases.code("Application"), cases.code("Module"))
    release = cases.code("Release")
    data = SlotCreation(namespace_id=namespace["id"], module_id=module["id"], release_name=release)
    key = {"namespace_id": namespace["id"], "release_name": release}
    i.merge(i.storage.slots, slots, key, data, f"{_namespace_name(cases)} / {release}")


def _installation(i: _ImportRun, cases: Row) -> None:
    namespace = i.namespace(cases)
    release = cases.code("Release")
    slot = i.find(
        i.storage.slots,
        f"emplacement « {_namespace_name(cases)} / {release} »",
        namespace_id=namespace["id"],
        release_name=release,
    )
    module = i.storage.modules.read_existing(slot["module_id"])
    version = i.version(module, cases.text("Version"))
    data = InstallationCreation(
        slot_id=slot["id"],
        module_version_id=version["id"],
        technical_parameter_id=i.parameter(module, "TECHNICAL", cases.code("Paramètre technique"))["id"],
        application_parameter_id=i.parameter(module, "APPLICATION", cases.code("Paramètre applicatif"))["id"],
        date=cases.date("Date (UTC)"),
    )
    key = data.model_dump()
    i.once(tuple(key.values()))
    described = f"{_namespace_name(cases)} / {release} : {version['number']} le {data.date:%Y-%m-%d %H:%M:%S}"
    fields = ("module_version_id", "technical_parameter_id", "application_parameter_id")
    already = any(
        all(existing_installation[field] == key[field] for field in fields)
        and to_utc(existing_installation["date"]) == data.date
        for existing_installation in i.storage.installations.list_all(slot_id=slot["id"])
    )
    if already:  # un fait passé : jamais modifié
        i.log(None, described)
        return
    installations.add(i.storage, key)
    i.log(ImportAction.ADD, described)


_PROCESSING: dict[str, Callable[[_ImportRun, Row], None]] = {
    "Environnements": _environment,
    "Clusters": _cluster,
    "Zones réseau": _network_zone,
    "Applications": _application,
    "Namespaces": _namespace,
    "Modules": _module,
    "Dépendances": _dependency,
    "Versions": _version,
    "Compatibilités": _compatibility,
    "Paramètres": _parameter,
    "Affectations": _assignment,
    "Emplacements": _slot,
    "Installations": _installation,
}
assert list(_PROCESSING) == list(SHEETS)  # un traitement par onglet, dans le même ordre
